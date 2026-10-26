#!/usr/bin/env python3
"""Analyze only the 'Crosslinker Co' sheet of Hana  python sheet.xlsx.

Jupyter, after changing to the folder containing this file and the workbook:
    %run crosslinker_co_kinetics_14.py

Terminal:
    python crosslinker_co_kinetics_14.py
    python crosslinker_co_kinetics_14.py --input /path/to/workbook.xlsx

Requires numpy, matplotlib, and openpyxl. Writes PNG plots, CSV tables, an
Excel workbook, and PDF and Markdown reports to the crosslinker_co_14_graphs folder.
Every header-named sample in this sheet is included automatically.
"""

import argparse
import csv
import math
import re
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


WORKING_DIR = Path.cwd()  # Jupyter's %cd location


def latest_workbook(directory):
    """Use the most recently updated matching workbook in Jupyter's folder."""
    names = re.compile(r"Hana  python sheet(?:\(\d+\))?\.xlsx")
    candidates = [path for path in directory.glob("Hana  python sheet*.xlsx")
                  if names.fullmatch(path.name)]
    return (max(candidates, key=lambda path: path.stat().st_mtime)
            if candidates else directory / "Hana  python sheet.xlsx")


DEFAULT_INPUT = latest_workbook(WORKING_DIR)
SHEET_NAME = 'Crosslinker Co'
OUTPUT_STEM = "crosslinker_co_14"
CHECK_METAL_LABELS = True
DEFAULT_OUTPUT = WORKING_DIR / f"{OUTPUT_STEM}_graphs"
MIN_POINTS = 6
MAX_INITIAL_OMISSIONS = 5  # Permit a short curved transient before the early fit.
MAX_TRANSITION_OMISSIONS = 2
MAX_BETWEEN_SAMPLE_SHIFT = 2  # Index difference at each of four window boundaries.
R2_THRESHOLD = 0.99
OMISSION_PENALTY_R2 = 0.001
GRAPH_TIME_LIMIT_MIN = 180
Y_AXIS_STEP = 0.05

# Optional display aliases for previously named experiments. Adding or removing
# a sample only requires changing the worksheet, not this mapping.
EXPERIMENT_LABELS = {}

MODELS = {
    "Fickian sqrt(t)": (r"$t^{1/2}$ (min$^{1/2}$)",
                        r"Metal concentration, $[M]$ (mol L$^{-1}$)", "mol L^-1 min^-1/2"),
    "First order": ("Time (min)", r"$\ln([M])$", "min^-1"),
    "AA": ("Time (min)", r"$1/[M]$ (L mol$^{-1}$)", "L mol^-1 min^-1"),
    "AB 1:1": ("Time (min)",
               r"1:1 AB linearization (L mol$^{-1}$)",
               "L mol^-1 min^-1"),
}

AXIS_LABEL_SIZE = 16
TICK_LABEL_SIZE = 12
LEGEND_LOCATION = "upper right"
LEGEND_ANCHOR = (0.98, 0.98)  # Axes fraction: identical for every plot.


def format_kinetic_axes(ax, xlabel, ylabel, x_limit, visible_values):
    """Apply the same axis typography, ranges, and legend to every PNG plot."""
    ax.set_xlabel(xlabel, fontsize=AXIS_LABEL_SIZE, labelpad=10)
    ax.set_ylabel(ylabel, fontsize=AXIS_LABEL_SIZE, labelpad=10)
    ax.tick_params(axis="both", labelsize=TICK_LABEL_SIZE)
    ax.set_xlim(0, x_limit)
    y_min, y_max = min(visible_values, default=0), max(visible_values, default=0)
    bottom = min(0, math.floor(y_min / Y_AXIS_STEP) * Y_AXIS_STEP)
    top = max(0, math.ceil(y_max / Y_AXIS_STEP) * Y_AXIS_STEP)
    if top == bottom:
        top += Y_AXIS_STEP
    ax.set_ylim(bottom, top)
    ax.grid(alpha=.2)
    ax.legend(loc=LEGEND_LOCATION, bbox_to_anchor=LEGEND_ANCHOR, ncol=2,
              fontsize=8, frameon=True, framealpha=.85, borderaxespad=0)


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def label_metal(name):
    """Return an unambiguous metal in a sample name, if it has one."""
    symbols = re.findall(r"(?<![A-Za-z])(Cu|Co|Ni|Mg|Ca|Fe|Zn)(?![A-Za-z])",
                         name, flags=re.IGNORECASE)
    unique = {symbol.casefold() for symbol in symbols}
    return next(iter(unique)) if len(unique) == 1 else None


def discover_groups(sheet):
    """Discover samples and their replicate columns from this sheet's headers."""
    headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column+1)]
    time_columns = []
    time_zero_rows, time_steps = {}, {}
    for col, value in enumerate(headers, 1):
        first_numeric = [(row, sheet.cell(row, col).value)
                         for row in range(2, min(sheet.max_row, 12)+1)
                         if is_number(sheet.cell(row, col).value)][:2]
        has_time_series = (len(first_numeric) == 2 and first_numeric[0][1] == 0 and
                           first_numeric[1][1] > 0)
        named_time = isinstance(value, str) and re.fullmatch(
            r"time\s*(?:/|\()?\s*min(?:utes?)?\s*\)?",
            " ".join(value.split()), re.IGNORECASE)
        if named_time or (value is None and has_time_series):
            time_columns.append(col)
            if has_time_series:
                time_zero_rows[col] = first_numeric[0][0]
                time_steps[col] = first_numeric[1][1] - first_numeric[0][1]
    if not time_columns:
        raise ValueError(f"No time-series columns found in worksheet {SHEET_NAME!r}")
    groups = {}
    time_col = None
    active_key = None
    for col, raw_header in enumerate(headers, 1):
        if col in time_columns:
            time_col, active_key = col, None
            continue
        if isinstance(raw_header, str) and raw_header.strip():
            if time_col is None:
                continue  # A note preceding the first time series is not a sample.
            zero_row = time_zero_rows.get(time_col)
            if (zero_row is None or
                    not is_number(sheet.cell(zero_row, col).value) or
                    sheet.cell(zero_row, col).value <= 0):
                active_key = None  # Ignore notes and empty columns in row 1.
                continue
            header = " ".join(raw_header.split())
            match = re.fullmatch(r"(.+?)\s+(?:rep(?:licate)?\s*)?(\d+)\s*(\([^)]*\))?",
                                 header, flags=re.IGNORECASE)
            base = (match.group(1).strip()+" "+match.group(3)
                    if match and match.group(3) else
                    match.group(1).strip() if match else header)
            base = " ".join(base.split())
            active_key = time_col, base.casefold()
            if active_key not in groups:
                step = f"{time_steps[time_col]:g}" if time_col in time_steps else ""
                label = EXPERIMENT_LABELS.get(f"{base.casefold()}|{step}",
                                              EXPERIMENT_LABELS.get(base.casefold(), base))
                groups[active_key] = dict(name=label, time_col=time_col, columns=[])
            groups[active_key]["columns"].append(col)
        elif (active_key is not None and raw_header is None and
              time_col in time_zero_rows and
              is_number(sheet.cell(time_zero_rows[time_col], col).value) and
              sheet.cell(time_zero_rows[time_col], col).value > 0):
            # A blank header directly after named replicates is still part of
            # that experiment when it contains concentration observations.
            groups[active_key]["columns"].append(col)
        else:
            active_key = None
    if not groups:
        raise ValueError(f"No named replicate columns found in worksheet {SHEET_NAME!r}")
    result, used_names = [], set()
    for item in groups.values():
        label = item["name"]
        if label in used_names:
            step = time_steps.get(item["time_col"])
            candidate = f"{label} ({step:g} min sampling)" if step else label
            if candidate in used_names:
                candidate = f"{label} (time column {get_column_letter(item['time_col'])})"
            label = candidate
        used_names.add(label)
        result.append((label, item["time_col"], tuple(item["columns"])))
    return result


def find_metadata_row(sheet, time_col, kind):
    """Find a group's own metal or crosslinker-mole row by its label."""
    metal_labels = {"cu", "co", "ni", "mg", "ca", "fe", "zn", "metal mole", "metal moles"}
    for row in range(2, sheet.max_row+1):
        value = sheet.cell(row, time_col).value
        if not isinstance(value, str):
            continue
        label = " ".join(value.casefold().split())
        if kind == "metal" and label in metal_labels:
            return row
        if kind == "linker" and label in {"crosslinker mole", "crosslinker moles"}:
            return row
    return None


def read_sheet(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[SHEET_NAME]  # The only worksheet accessed.
        experiments = []
        for name, time_col, concentration_cols in discover_groups(sheet):
            metal_row = find_metadata_row(sheet, time_col, "metal")
            linker_row = find_metadata_row(sheet, time_col, "linker")
            source_metal_label = sheet.cell(metal_row, time_col).value if metal_row else None
            def mole_values(row):
                values = [sheet.cell(row, col).value if row else None
                          for col in concentration_cols]
                return np.array([float(v) if is_number(v) else np.nan for v in values]), values

            metal, raw_metal = mole_values(metal_row)
            linker, raw_linker = mole_values(linker_row)
            issues = []
            for kind, row, raw, values, valid in (
                ("metal", metal_row, raw_metal, metal, np.isfinite(metal) & (metal > 0)),
                ("crosslinker", linker_row, raw_linker, linker,
                 np.isfinite(linker) & (linker >= 0))):
                if row is None:
                    issues.append(f"{kind}-mole row label missing in {get_column_letter(time_col)}")
                else:
                    issues.extend(f"{kind} moles {get_column_letter(col)}{row}={value!r}"
                                  for col, value, good in zip(concentration_cols, raw, valid)
                                  if not good)
            times, concentrations = [], []
            data_stop = min((row for row in (metal_row, linker_row) if row is not None),
                            default=sheet.max_row+1)
            for row in range(2, data_stop):
                t = sheet.cell(row, time_col).value
                values = [sheet.cell(row, c).value for c in concentration_cols]
                if is_number(t) and all(is_number(v) for v in values):
                    times.append(float(t))
                    concentrations.append([float(v) for v in values])
            t, C = np.array(times), np.array(concentrations)
            if len(t) < 2 * MIN_POINTS or t[0] != 0 or np.any(np.diff(t) <= 0) or np.any(C <= 0):
                raise ValueError(f"Invalid time/concentration series: {name}")
            experiments.append(dict(name=name, time=t, replicates=C,
                                    metal=metal, linker=linker,
                                    ab_metadata_reason="; ".join(issues),
                                    source_metal_label=source_metal_label))
        return experiments
    finally:
        book.close()


def transformed_series(experiment):
    """Return x/y series and an AB unavailability reason, if any."""
    t, C = experiment["time"], experiment["replicates"]
    mean_C = C.mean(axis=1)
    series = {
        "Untreated": (t, mean_C),
        "Fickian sqrt(t)": (np.sqrt(t), mean_C),
        "First order": (t, np.log(mean_C)),
        "AA": (t, 1 / mean_C),
    }
    if experiment["ab_metadata_reason"]:
        return series, experiment["ab_metadata_reason"]
    if np.any(experiment["linker"] == 0):
        return series, "initial crosslinker moles recorded as zero"

    # For each replicate infer V = n(metal,0)/C(metal,0). With 1:1 consumption,
    # B0 = A0*n(crosslinker,0)/n(metal,0), B(t) = B0 - (A0-A(t)).
    A0 = C[0, :]
    B0 = A0 * experiment["linker"] / experiment["metal"]
    delta = B0 - A0
    B = C + delta[None, :]
    invalid = np.flatnonzero(np.any(B <= 0, axis=1))
    valid_n = int(invalid[0]) if len(invalid) else len(t)
    if valid_n < 2 * MIN_POINTS:
        return series, "fewer than twelve consecutive time points with positive crosslinker"
    ab = np.empty_like(C[:valid_n])
    for j in range(C.shape[1]):
        if abs(delta[j]) < 1e-8 * A0[j]:
            ab[:, j] = 1 / C[:valid_n, j] - 1 / A0[j]
        else:
            ab[:, j] = (np.log1p(delta[j] / C[:valid_n, j])
                        - np.log1p(delta[j] / A0[j])) / delta[j]
    series["AB 1:1"] = (t[:valid_n], ab.mean(axis=1))
    return series, ""


def line_fit(x, y):
    slope, intercept = np.polyfit(x, y, 1)
    sse = float(np.sum((y - (slope*x + intercept))**2))
    sst = float(np.sum((y - np.mean(y))**2))
    # A constant plateau has no explained concentration variation.
    r2 = 1 - sse/sst if sst > 0 else 0.0
    return dict(slope=float(slope), intercept=float(intercept), sse=sse, r2=float(r2))


def apparent_k(model, slope):
    return -slope if model in ("Untreated", "Fickian sqrt(t)", "First order") else slope


def forward_maximum(grid):
    """For each index, find the best candidate in its forward 3-point box."""
    current = grid
    for axis in range(grid.ndim):
        following = current.copy()
        for distance in range(1, MAX_BETWEEN_SAMPLE_SHIFT + 1):
            source = [slice(None)] * grid.ndim
            target = [slice(None)] * grid.ndim
            source[axis] = slice(distance, None)
            target[axis] = slice(None, -distance)
            np.maximum(following[tuple(target)], current[tuple(source)],
                       out=following[tuple(target)])
        current = following
    return current


def optimize_group(experiments, series_by_name, model_options, window_caches,
                   diagnostic=False):
    """Search all samples together, then choose their best compatible windows.

    A common four-coordinate box is searched globally. Within a box each
    sample chooses its best model/window independently, so no sample acts as
    the reference fit. Both endpoints of both segments differ by at most two
    observation indices between any two samples. A missing AB model simply
    excludes that model for the affected sample.
    """
    eligible = [exp for exp in experiments if any(
        first in series_by_name[exp["name"]] and second in series_by_name[exp["name"]]
        for first, second in model_options)]
    if not eligible:
        return {}
    # Longer series may use at most two points beyond the shortest series.
    max_stop = min(len(exp["time"]) for exp in eligible) + MAX_BETWEEN_SAMPLE_SHIFT
    shape = (MAX_INITIAL_OMISSIONS + 1,) + (max_stop + 1,) * 3
    grids = {}
    for exp in eligible:
        name = exp["name"]
        series = series_by_name[name]
        cache = window_caches[name]
        passing = np.full(shape, -np.inf)
        fallback = np.full(shape, -np.inf)
        pass_option = np.full(shape, -1, dtype=np.int16)
        fallback_option = np.full(shape, -1, dtype=np.int16)

        def window(model, start, stop):
            key = model, start, stop
            if key not in cache:
                x, y = series[model]
                cache[key] = line_fit(x[start:stop], y[start:stop])
            return cache[key]

        for option_id, (early_model, late_model) in enumerate(model_options):
            if early_model not in series or late_model not in series:
                continue
            if not diagnostic and (late_model == "Fickian sqrt(t)" or
                                   early_model == late_model or
                                   {early_model, late_model} <= {"AA", "AB 1:1"}):
                continue
            early_length = len(series[early_model][0])
            late_length = len(series[late_model][0])
            last_stop = min(late_length, max_stop)
            max_initial = min(MAX_INITIAL_OMISSIONS, early_length - MIN_POINTS,
                              last_stop - 2 * MIN_POINTS)
            for initial in range(max_initial + 1):
                for split in range(initial + MIN_POINTS,
                                   min(early_length, last_stop - MIN_POINTS) + 1):
                    early = window(early_model, initial, split)
                    for gap in range(MAX_TRANSITION_OMISSIONS + 1):
                        start = split + gap
                        if start + MIN_POINTS > last_stop:
                            break
                        for stop in range(start + MIN_POINTS, last_stop + 1):
                            late = window(late_model, start, stop)
                            quality = ((early["r2"] + late["r2"]) / 2 -
                                       OMISSION_PENALTY_R2 *
                                       (initial + gap + late_length - stop))
                            both_pass = (early["r2"] > R2_THRESHOLD and
                                         late["r2"] > R2_THRESHOLD)
                            index = initial, split, start, stop
                            values = passing if both_pass else fallback
                            choices = pass_option if both_pass else fallback_option
                            if quality > values[index]:
                                values[index] = quality
                                choices[index] = option_id
        if not np.isfinite(passing).any() and not np.isfinite(fallback).any():
            raise ValueError(f"No valid two-window fit for {name}")
        grids[name] = (passing, fallback, pass_option, fallback_option)

    available = np.ones(shape, dtype=bool)
    passing_count = np.zeros(shape, dtype=np.int16)
    total_quality = np.zeros(shape)
    for exp in eligible:
        passing, fallback, _, _ = grids[exp["name"]]
        best_pass = forward_maximum(passing)
        best_fallback = forward_maximum(fallback)
        has_pass = np.isfinite(best_pass)
        has_fit = has_pass | np.isfinite(best_fallback)
        available &= has_fit
        passing_count += has_pass
        total_quality += np.where(has_pass, best_pass,
                                  np.where(has_fit, best_fallback, 0))
    if not available.any():
        raise ValueError("No compatible windows: the two-point sample limit "
                         "conflicts with six-point segments")
    highest_pass_count = int(passing_count[available].max())
    contenders = np.flatnonzero(available & (passing_count == highest_pass_count))
    anchor = np.unravel_index(contenders[np.argmax(total_quality.flat[contenders])], shape)

    selected = {}
    neighborhood = tuple(slice(position, position + MAX_BETWEEN_SAMPLE_SHIFT + 1)
                         for position in anchor)
    for exp in eligible:
        name = exp["name"]
        passing, fallback, pass_option, fallback_option = grids[name]
        use_pass = np.isfinite(passing[neighborhood]).any()
        values = passing if use_pass else fallback
        choices = pass_option if use_pass else fallback_option
        neighborhood_values = values[neighborhood]
        relative = np.unravel_index(np.argmax(neighborhood_values),
                                    neighborhood_values.shape)
        initial, split, start, stop = (anchor[axis] + relative[axis]
                                       for axis in range(4))
        option_id = int(choices[initial, split, start, stop])
        early_model, late_model = model_options[option_id]
        early = window_caches[name][early_model, initial, split]
        late = window_caches[name][late_model, start, stop]
        selected[name] = (early_model, late_model, initial, split,
                          start, stop, early, late)
    for boundary in (2, 3, 4, 5):
        indices = [choice[boundary] for choice in selected.values()]
        assert max(indices) - min(indices) <= MAX_BETWEEN_SAMPLE_SHIFT
    return selected


def pair_from_choice(experiment, series, choice):
    early_model, late_model, initial, split, start, stop, early, late = choice
    gap = start - split
    t = experiment["time"]
    early_end, late_start = t[split-1], t[split+gap]
    assert early_end < late_start
    late_length = len(series[late_model][0])
    late_n = stop-split-gap
    return dict(experiment=experiment["name"], early_model=early_model,
                early_from=t[initial], early_to=early_end, early_n=split-initial,
                early_r2=early["r2"], early_k=apparent_k(early_model, early["slope"]),
                early_unit=MODELS[early_model][2], late_model=late_model,
                late_from=late_start, late_to=t[stop-1], late_n=late_n,
                late_r2=late["r2"], late_k=apparent_k(late_model, late["slope"]),
                late_unit=MODELS[late_model][2], early_gap=gap, late_gap=gap,
                initial_omitted_points=initial,
                initial_omitted_times=";".join(f"{v:g}" for v in t[:initial]),
                omitted_times=";".join(f"{v:g}" for v in t[split:split+gap]),
                intervening_points=gap,
                intervening_times=";".join(f"{v:g}" for v in t[split:split+gap]),
                late_unused_tail_points=late_length-stop,
                late_unused_tail_times=";".join(f"{v:g}" for v in t[stop:late_length]),
                segment_relation=(f"separate; early ends {early_end:g}, "
                                  f"late starts {late_start:g} min; {gap} intervening points"),
                weighted_mean_r2=((split-initial)*early["r2"]+late_n*late["r2"])/
                                 (split-initial+late_n),
                both_above_099=(early["r2"] > R2_THRESHOLD and
                                late["r2"] > R2_THRESHOLD),
                original_early_model=early_model,
                adjustment_reason="optimized across all samples")


def write_csv(path, columns, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def second_order_rows(experiments, series_by_name, ab_reasons, early_model, window_caches):
    """Jointly optimize each comparison across all eligible samples."""
    joint_choices = {
        model: optimize_group(experiments, series_by_name,
                              ((early_model, model),), window_caches)
        for model in ("AA", "AB 1:1")
    }
    rows = []
    for experiment in experiments:
        name = experiment["name"]
        options = {}
        for model in ("AA", "AB 1:1"):
            choice = joint_choices[model].get(name)
            if choice is None:
                options[model] = dict(experiment=name, early_model=early_model,
                                      late_model=model, preferred_second_order=False,
                                      status=f"unavailable: {ab_reasons.get(name, 'insufficient valid points')}")
            else:
                result = pair_from_choice(experiment, series_by_name[name], choice)
                options[model] = dict(result, preferred_second_order=False,
                                      status="fitted")
        available = [m for m, row in options.items() if row["status"] == "fitted"]
        if available:
            preferred = max(available, key=lambda m: (
                options[m]["both_above_099"],
                (options[m]["early_r2"]+options[m]["late_r2"])/2 -
                OMISSION_PENALTY_R2*(options[m]["initial_omitted_points"]+
                                     options[m]["early_gap"]+
                                     options[m]["late_unused_tail_points"])))
            options[preferred]["preferred_second_order"] = True
        rows.extend(options[model] for model in ("AA", "AB 1:1"))
    return rows


def write_excel_results(path, fit_rows, pairs, fickian_rows, first_order_rows):
    """Put both restricted comparisons alongside the source fit tables."""
    book = Workbook()
    def summary_columns(early_label, early_unit):
        return [
        ("experiment", "Experiment"),
        ("late_model", "Second-order model"),
        ("preferred_second_order", "Preferred"),
        ("early_from", f"{early_label} from (min)"),
        ("early_to", f"{early_label} to (min)"),
        ("early_n", f"{early_label} points"),
        ("early_r2", f"{early_label} R²"),
        ("early_k", f"{early_label} k ({early_unit})"),
        ("initial_omitted_points", "Initial transient (points)"),
        ("initial_omitted_times", "Initial transient times (min)"),
        ("late_from", "Late from (min)"),
        ("late_to", "Late to (min)"),
        ("late_n", "Late points"),
        ("late_r2", "Late R²"),
        ("late_k", "Late k (L/mol/min)"),
        ("late_unused_tail_points", "Unused tail points"),
        ("late_unused_tail_times", "Unused tail times (min)"),
        ("intervening_points", "Between windows (points)"),
        ("segment_relation", "Time relation"),
        ("both_above_099", "Both R² > 0.99"),
        ("status", "Status"),
        ]
    sheets = [
        ("Fickian + 2nd order", fickian_rows),
        ("First order + 2nd order", first_order_rows),
        ("Best pairs", pairs),
        ("Model fits", fit_rows),
    ]
    for index, (title, rows) in enumerate(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = title
        columns = (summary_columns("Fickian", "mol/L/min^0.5") if index == 0 else
                   summary_columns("First order", "min^-1") if index == 1 else
                   [(key, key) for key in dict.fromkeys(key for row in rows for key in row)])
        sheet.append([label for _, label in columns])
        for row in rows:
            sheet.append([row.get(key) for key, _ in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17324A")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for col, (key, label) in enumerate(columns, 1):
            text_width = max(len(label), *(len(str(row.get(key, ""))) for row in rows))
            sheet.column_dimensions[get_column_letter(col)].width = min(max(text_width+2, 13), 36)
        for row in sheet.iter_rows(min_row=2):
            if index < 2 and row[2].value is True:
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor="E7F3EA")
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
                if index < 2:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
            if index < 2:
                sheet.row_dimensions[row[0].row].height = 30
    book.save(path)


def draw_plot(model, experiments, series_by_name, fits, output, full_range=False):
    fig, ax = plt.subplots(figsize=(8, 6), constrained_layout=True)
    colors = plt.get_cmap("tab10").colors
    max_time = max(exp["time"][-1] for exp in experiments)
    displayed_time = max_time if full_range else GRAPH_TIME_LIMIT_MIN
    x_limit = math.sqrt(displayed_time) if model == "Fickian sqrt(t)" else displayed_time
    visible_values = []
    for index, experiment in enumerate(experiments):
        name = experiment["name"]
        if model not in series_by_name[name]:
            continue
        x, y = series_by_name[name][model]
        color = colors[index]
        ax.plot(x, y, "-o", color=color, markersize=8, lw=.8, alpha=.65, label=name)
        visible_values.extend(y[x <= x_limit + 1e-10])
        if model != "Untreated":
            full, split, gap, stop, early, late = fits[(name, model)]
            for sub_x, result in ((x[early["start"]:split], early),
                                  (x[split+gap:stop], late)):
                if result["r2"] > R2_THRESHOLD:
                    line_y = result["slope"]*sub_x+result["intercept"]
                    ax.plot(sub_x, line_y, color=color, lw=1.7)
                    visible_values.extend(line_y[sub_x <= x_limit + 1e-10])
    xlabel, ylabel, _ = MODELS.get(
        model, ("Time (min)", r"Metal concentration, $[M]$ (mol L$^{-1}$)", ""))
    format_kinetic_axes(ax, xlabel, ylabel, x_limit, visible_values)
    filename = model.lower().replace(" ", "_").replace(":", "").replace("/", "_")
    fig.savefig(output / (filename + ("_full" if full_range else "") + ".png"), dpi=190)
    return fig


def draw_first_order_aa_plot(phase, experiments, series_by_name, first_order_rows,
                             output, full_range=False):
    """Plot the actual jointly selected First-order + AA window for each group."""
    assert phase in ("early", "late")
    model = "First order" if phase == "early" else "AA"
    chosen = {row["experiment"]: row for row in first_order_rows
              if row["late_model"] == "AA" and row["status"] == "fitted"}
    fig, ax = plt.subplots(figsize=(8, 6), constrained_layout=True)
    colors = plt.get_cmap("tab10").colors
    max_time = max(exp["time"][-1] for exp in experiments)
    x_limit = max_time if full_range else GRAPH_TIME_LIMIT_MIN
    visible_values = []
    for index, experiment in enumerate(experiments):
        name = experiment["name"]
        if name not in chosen:
            continue
        row = chosen[name]
        t = experiment["time"]
        x, y = series_by_name[name][model]
        start = int(np.searchsorted(t, row[f"{phase}_from"]))
        stop = int(np.searchsorted(t, row[f"{phase}_to"], side="right"))
        assert stop - start == row[f"{phase}_n"] >= MIN_POINTS
        assert t[start] == row[f"{phase}_from"] and t[stop-1] == row[f"{phase}_to"]
        result = line_fit(x[start:stop], y[start:stop])
        assert math.isclose(result["r2"], row[f"{phase}_r2"], abs_tol=1e-9)
        color = colors[index % len(colors)]
        # Every measured mean is shown; the line stops at the selected window.
        ax.plot(x, y, "-o", color=color, markersize=8, lw=.8, alpha=.65,
                label=f"{name} (R²={result['r2']:.3f})")
        line_x = np.linspace(x[start], x[stop-1], 80)
        line_y = result["slope"]*line_x + result["intercept"]
        ax.plot(line_x, line_y, color=color, lw=1.8)
        visible_values.extend(y[x <= x_limit + 1e-10])
        visible_values.extend(line_y[line_x <= x_limit + 1e-10])
    xlabel, ylabel, _ = MODELS[model]
    format_kinetic_axes(ax, xlabel, ylabel, x_limit, visible_values)
    filename = "first_order_early_aa_pair" if phase == "early" else "aa_late_first_order_pair"
    fig.savefig(output / (filename + ("_full" if full_range else "") + ".png"), dpi=190)
    return fig


def display_number(value):
    return f"{value:.6g}" if isinstance(value, (int, float, np.floating)) else str(value)


def report_table(pdf, title, subtitle, columns, widths, rows, rows_per_page=18):
    for offset in range(0, len(rows), rows_per_page):
        block = rows[offset:offset+rows_per_page]
        fig, ax = plt.subplots(figsize=(11.7, 8.3))
        ax.axis("off")
        fig.text(.055, .945, title, fontsize=17, weight="bold", color="#17324a")
        fig.text(.055, .91, subtitle, fontsize=9, color="#4b5966")
        short = len(block) < 9
        table = ax.table(cellText=block, colLabels=columns, cellLoc="left", colLoc="left",
                         colWidths=widths,
                         bbox=[.015, .42 if short else .06, .97, .44 if short else .8])
        table.auto_set_font_size(False)
        table.set_fontsize(8 if short else 7.2)
        for (i, _), cell in table.get_celld().items():
            cell.set_linewidth(.3)
            cell.set_edgecolor("#dbe3e8")
            cell.set_facecolor("#17324a" if i == 0 else ("#f0f5f7" if i % 2 == 0 else "white"))
            if i == 0:
                cell.get_text().set_color("white")
                cell.get_text().set_weight("bold")
        fig.text(.055, .03, f"Rows {offset+1}-{offset+len(block)} of {len(rows)}",
                 fontsize=8, color="#667781")
        pdf.savefig(fig)
        plt.close(fig)


def write_pdf_report(path, experiments, series_by_name, fits, fit_rows, pairs,
                     fickian_rows, first_order_rows, ab_reasons, label_notes, output):
    with PdfPages(path) as pdf:
        fig = plt.figure(figsize=(11.7, 8.3))
        fig.text(.07, .88, f"{SHEET_NAME} | concentration kinetics", fontsize=24,
                 weight="bold", color="#17324a")
        notes = [
            f"Input: the {SHEET_NAME} worksheet only; source values are assumed mol/L.",
            "All available experiments appear together; connected 8-point markers show replicate means.",
            "For 1:1 consumption, B(t)=B0-[A0-A(t)]; AB y=[ln(B/A)-ln(B0/A0)]/(B0-A0).",
            "Untreated: points only. Selected diagnostic window lines require R² > 0.99.",
            "First-order + AA pair plots show both selected fits; their R² appears in the legend.",
            "Graphs have no titles; [M] denotes metal concentration (mol/L).",
            "Legends use a fixed upper-right position inside every kinetic graph.",
            "Time views stop at 180 min; full-time views are added when needed.",
            "Exactly two fits use >=6 consecutive points each; an initial transient may be excluded.",
            "All samples jointly choose windows whose early and late boundaries differ by at most two observation positions.",
            "All model pairs refit disjoint windows; Fickian is excluded from late choices.",
            "High R² after a transformation alone does not establish a physical mechanism.",
        ]
        if label_notes:
            notes.append("Source metal labels differ from group names; numeric moles are used as recorded.")
        for i, note in enumerate(notes):
            fig.text(.07, .75-.044*i, note, fontsize=9.6)
        fig.text(.07, .145, "AB availability", fontsize=13, weight="bold", color="#9b522d")
        ab_lines = ([f"{name}: {reason}" for name, reason in ab_reasons.items()]
                    if ab_reasons else ["All groups have positive recorded crosslinker moles."])
        for i, line in enumerate(ab_lines):
            fig.text(.07, .105-.028*i, line, fontsize=9)
        pdf.savefig(fig)
        plt.close(fig)

        for model in ["Untreated", *MODELS]:
            if any(model in series_by_name[e["name"]] for e in experiments):
                fig = draw_plot(model, experiments, series_by_name, fits, output)
                pdf.savefig(fig)
                plt.close(fig)
                if any(e["time"][-1] > GRAPH_TIME_LIMIT_MIN for e in experiments):
                    fig = draw_plot(model, experiments, series_by_name, fits, output, full_range=True)
                    pdf.savefig(fig)
                    plt.close(fig)

        for phase in ("early", "late"):
            fig = draw_first_order_aa_plot(phase, experiments, series_by_name,
                                            first_order_rows, output)
            pdf.savefig(fig)
            plt.close(fig)
            if any(e["time"][-1] > GRAPH_TIME_LIMIT_MIN for e in experiments):
                fig = draw_first_order_aa_plot(phase, experiments, series_by_name,
                                                first_order_rows, output, full_range=True)
                pdf.savefig(fig)
                plt.close(fig)

        best_rows = [[p["experiment"], p["early_model"], p["late_model"],
                      f"{p['early_r2']:.4f}", f"{p['late_r2']:.4f}",
                      f"{p['weighted_mean_r2']:.4f}",
                      "yes" if p["both_above_099"] else "no",
                      p["adjustment_reason"] or "none"] for p in pairs]
        report_table(pdf, "Best early/late model pairs",
                     "Selected R² values and time ranges exactly match rows in model_fits.csv.",
                     ["Experiment", "Early model", "Late model", "Early R²", "Late R²",
                      "Mean R²", "Both >.99", "Adjustment"],
                     [.15, .16, .16, .08, .08, .08, .09, .18], best_rows)

        def summary_rows(source_rows):
            return [
            [r["experiment"],
             f"{r['early_from']:g}-{r['early_to']:g}" if r["status"] == "fitted" else "n/a",
             f"{r['early_r2']:.4f}" if r["status"] == "fitted" else "n/a", r["late_model"],
             f"{r['late_from']:g}-{r['late_to']:g}" if r["status"] == "fitted" else "n/a",
             f"{r['late_r2']:.4f}" if r["status"] == "fitted" else "n/a",
             "yes" if r["preferred_second_order"] else "no", r["status"]]
            for r in source_rows]
        report_table(pdf, "Fickian early + AA or AB late",
                     "Both windows are jointly optimized; preferred means higher combined fit score.",
                     ["Experiment", "Fickian time", "Early R²", "Late", "Late time",
                      "Late R²", "Preferred", "Status"],
                     [.15, .12, .09, .10, .12, .09, .10, .23], summary_rows(fickian_rows))
        report_table(pdf, "First order early + AA or AB late",
                     "Both windows are jointly optimized with at most two transition points omitted.",
                     ["Experiment", "1st-order time", "Early R²", "Late", "Late time",
                      "Late R²", "Preferred", "Status"],
                     [.15, .12, .09, .10, .12, .09, .10, .23], summary_rows(first_order_rows))

        for offset in range(0, len(pairs), 3):
            fig = plt.figure(figsize=(11.7, 8.3))
            fig.text(.055, .95, "Selected pair | experiment detail", fontsize=17,
                     weight="bold", color="#17324a")
            for j, p in enumerate(pairs[offset:offset+3]):
                top = .85-j*.27
                fig.text(.07, top, p["experiment"], fontsize=12, weight="bold", color="#17324a")
                ax = fig.add_axes([.07, top-.16, .86, .12])
                ax.axis("off")
                rows = [[phase.title(), p[f"{phase}_model"],
                         f"{p[f'{phase}_from']:g}-{p[f'{phase}_to']:g}",
                         f"{p[f'{phase}_r2']:.5f}", f"{p[f'{phase}_k']:.6g}",
                         p[f"{phase}_unit"]] for phase in ("early", "late")]
                tab = ax.table(cellText=rows,
                               colLabels=["Segment", "Model", "Time (min)", "R²", "k_app", "k units"],
                               colWidths=[.12, .23, .13, .10, .13, .29], loc="center", cellLoc="left")
                tab.auto_set_font_size(False)
                tab.set_fontsize(8)
                tab.scale(1, 1.3)
                for (i, _), cell in tab.get_celld().items():
                    cell.set_linewidth(.3)
                    cell.set_edgecolor("#dbe3e8")
                    cell.set_facecolor("#17324a" if i == 0 else ("#f0f5f7" if i == 2 else "white"))
                    if i == 0:
                        cell.get_text().set_color("white")
                fig.text(.075, top-.18,
                         f"Weighted mean R² {p['weighted_mean_r2']:.5f} | "
                         f"both >0.99: {'yes' if p['both_above_099'] else 'no'} | "
                         f"initial omitted: {p['initial_omitted_points']} | "
                         f"transition: {p['omitted_times'] or 'none'} | "
                         f"late unused tail: {p['late_unused_tail_points']} points",
                         fontsize=8, color="#4b5966")
                fig.text(.075, top-.205, p["segment_relation"], fontsize=8, color="#4b5966")
            pdf.savefig(fig)
            plt.close(fig)

        whole_rows = []
        segment_rows = []
        for e in experiments:
            name = e["name"]
            for model in MODELS:
                if (name, model) in fits:
                    full, _, _, _, early, late = fits[(name, model)]
                    whole_rows.append([name, model, str(len(series_by_name[name][model][0])),
                                       f"{full['r2']:.5f}", f"{apparent_k(model, full['slope']):.6g}",
                                       "yes" if full["r2"] < R2_THRESHOLD else "exploratory"])
                else:
                    whole_rows.append([name, model, "", "n/a", "n/a", "AB unavailable"])
        for row in fit_rows:
            if row.get("fit_scope") != "model_diagnostic":
                continue
            segment_rows.append([row["experiment"], row["model"], row["segment"],
                                 f"{row['time_from_min']}-{row['time_to_min']}" if row["status"] == "fitted" else "",
                                 str(row["n_points"]),
                                 f"{row['segment_r2']:.5f}" if row["status"] == "fitted" else "n/a",
                                 f"{row['apparent_k']:.6g}" if row["status"] == "fitted" else "n/a",
                                 row["k_unit"], str(row["unused_tail_points"])])
        report_table(pdf, "Whole-series model overview",
                     "Whole-series R² is diagnostic; two constrained windows are fitted for every available model.",
                     ["Experiment", "Model", "N", "Full R²", "Full k_app", "Split"],
                     [.22, .23, .055, .105, .14, .15], whole_rows)
        report_table(pdf, "Segment fits | four models",
                     "All model x early/late rows are shown; undefined AB transforms are labeled n/a.",
                     ["Experiment", "Model", "Segment", "Time (min)", "N", "R²", "k_app", "k units", "Tail N"],
                     [.15, .17, .08, .115, .035, .08, .10, .155, .075], segment_rows,
                     rows_per_page=12)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if not args.input.is_file():
        parser.error(f"Workbook not found: {args.input}. Use %cd to its folder or pass --input PATH.")
    args.out.mkdir(parents=True, exist_ok=True)
    print(f"Reading {args.input.resolve()} [worksheet: {SHEET_NAME}]")
    experiments = read_sheet(args.input)
    label_notes = [f"{e['name']}: metal-mole row says {e['source_metal_label']!r} "
                   f"(group header indicates {label_metal(e['name'])})"
                   for e in experiments
                   if CHECK_METAL_LABELS and label_metal(e["name"]) and
                   str(e["source_metal_label"]).casefold() != label_metal(e["name"])]
    series_by_name, ab_reasons, fits, fit_rows, window_caches = {}, {}, {}, [], {}
    untreated_rows = []
    for exp in experiments:
        name, t = exp["name"], exp["time"]
        series, reason = transformed_series(exp)
        series_by_name[name] = series
        window_caches[name] = {}
        untreated_rows.extend(dict(experiment=name, time_min=float(time),
                                   mean_concentration=float(conc),
                                   replicates=exp["replicates"].shape[1])
                              for time, conc in zip(*series["Untreated"]))
        if reason:
            ab_reasons[name] = reason

    # Select every model's diagnostic windows across all eligible samples.
    # All four boundaries (early start/end and late start/end) are compared.
    for model in MODELS:
        selections = optimize_group(experiments, series_by_name,
                                    ((model, model),), window_caches, diagnostic=True)
        for exp in experiments:
            name, t = exp["name"], exp["time"]
            series = series_by_name[name]
            if name in selections:
                _, _, initial, split, start, stop, early, late = selections[name]
                x, y = series[model]
                full = line_fit(x, y)
                gap = start - split
                fits[(name, model)] = full, split, gap, stop, dict(early, start=initial), late
                used_t = t[:len(x)]
                unit = MODELS[model][2]
                for segment, result, beginning, end in (
                    ("early", early, initial, split), ("late", late, start, stop)
                ):
                    fit_rows.append(dict(
                        experiment=name, model=model, segment=segment,
                        fit_scope="model_diagnostic", pair_id="", paired_model="",
                        status="fitted", split_required=full["r2"] < R2_THRESHOLD,
                        omitted_times_min=";".join(f"{v:g}" for v in used_t[split:start]),
                        initial_omitted_points=initial,
                        initial_omitted_times_min=";".join(f"{v:g}" for v in used_t[:initial]),
                        unused_tail_points=len(x)-stop,
                        unused_tail_times_min=";".join(f"{v:g}" for v in used_t[stop:]),
                        n_points=end-beginning, time_from_min=f"{used_t[beginning]:g}",
                        time_to_min=f"{used_t[end-1]:g}", full_r2=full["r2"],
                        segment_r2=result["r2"], apparent_k=apparent_k(model, result["slope"]),
                        k_unit=unit, full_apparent_k=apparent_k(model, full["slope"])))
                continue
            if model != "AB 1:1":
                raise ValueError(f"No valid diagnostic windows for {name}: {model}")
            for segment in ("early", "late"):
                fit_rows.append(dict(experiment=name, model=model, segment=segment,
                                     fit_scope="model_diagnostic", pair_id="", paired_model="",
                                     status=f"unavailable: {ab_reasons.get(name, 'insufficient valid points')}",
                                     split_required="", omitted_times_min="",
                                     initial_omitted_points="", initial_omitted_times_min="",
                                     unused_tail_points="", unused_tail_times_min="", n_points="",
                                     time_from_min="", time_to_min="", full_r2="",
                                     segment_r2="", apparent_k="", k_unit=MODELS[model][2],
                                     full_apparent_k=""))

    options = tuple((early, late) for early in MODELS for late in MODELS
                    if late != "Fickian sqrt(t)" and early != late and
                    not {early, late} <= {"AA", "AB 1:1"})
    chosen = optimize_group(experiments, series_by_name, options, window_caches)
    if len(chosen) != len(experiments):
        raise ValueError("One or more experiments have no valid mechanism pair")
    pairs = [pair_from_choice(exp, series_by_name[exp["name"]], chosen[exp["name"]])
             for exp in experiments]
    fickian_rows = second_order_rows(
        experiments, series_by_name, ab_reasons, "Fickian sqrt(t)", window_caches)
    first_order_rows = second_order_rows(
        experiments, series_by_name, ab_reasons, "First order", window_caches)

    # The diagnostic model rows above keep their model-specific two-window fits.
    # These additional rows reproduce the exact fits in every selected pair.
    for scope, source in (("best_pair", pairs),
                          ("fickian_second_order", fickian_rows),
                          ("first_order_second_order", first_order_rows)):
        for pair in source:
            if pair.get("status", "fitted") != "fitted":
                continue
            pair_id = f"{scope}::{pair['experiment']}::{pair['early_model']}::{pair['late_model']}"
            for phase, other in (("early", "late"), ("late", "early")):
                model = pair[f"{phase}_model"]
                full = fits[(pair["experiment"], model)][0]
                fit_rows.append(dict(
                    experiment=pair["experiment"], model=model, segment=phase,
                    fit_scope=scope, pair_id=pair_id, paired_model=pair[f"{other}_model"],
                    status="fitted", split_required=full["r2"] < R2_THRESHOLD,
                    omitted_times_min=pair["omitted_times"],
                    initial_omitted_points=pair["initial_omitted_points"],
                    initial_omitted_times_min=pair["initial_omitted_times"],
                    unused_tail_points=(pair["late_unused_tail_points"] if phase == "late" else 0),
                    unused_tail_times_min=(pair["late_unused_tail_times"] if phase == "late" else ""),
                    n_points=pair[f"{phase}_n"], time_from_min=f"{pair[f'{phase}_from']:g}",
                    time_to_min=f"{pair[f'{phase}_to']:g}", full_r2=full["r2"],
                    segment_r2=pair[f"{phase}_r2"], apparent_k=pair[f"{phase}_k"],
                    k_unit=pair[f"{phase}_unit"],
                    full_apparent_k=apparent_k(model, full["slope"])))

    fit_cols = ["experiment", "model", "segment", "fit_scope", "pair_id", "paired_model",
                "status", "split_required",
                "omitted_times_min", "initial_omitted_points", "initial_omitted_times_min",
                "unused_tail_points", "unused_tail_times_min",
                "n_points", "time_from_min", "time_to_min",
                "full_r2", "segment_r2", "apparent_k", "k_unit", "full_apparent_k"]
    write_csv(args.out / "model_fits.csv", fit_cols, fit_rows)
    write_csv(args.out / "untreated_means.csv",
              ["experiment", "time_min", "mean_concentration", "replicates"], untreated_rows)
    pair_cols = list(pairs[0].keys())
    write_csv(args.out / "best_model_pairs.csv", pair_cols, pairs)
    summary_cols = list(dict.fromkeys(
        key for row in (*fickian_rows, *first_order_rows) for key in row))
    write_csv(args.out / "fickian_second_order_pairs.csv", summary_cols, fickian_rows)
    write_csv(args.out / "first_order_second_order_pairs.csv", summary_cols, first_order_rows)
    write_excel_results(args.out / f"{OUTPUT_STEM}_kinetics_results.xlsx",
                        fit_rows, pairs, fickian_rows, first_order_rows)
    write_pdf_report(args.out / f"{OUTPUT_STEM}_report.pdf", experiments,
                     series_by_name, fits, fit_rows, pairs, fickian_rows, first_order_rows,
                     ab_reasons, label_notes, args.out)

    lines = [f"# {SHEET_NAME} kinetics", "",
             f"Input worksheet: `{SHEET_NAME}` only. Time is in minutes; concentrations are assumed mol/L.",
             "Every sample is discovered from the row-1 headers and its adjacent replicate columns. Adding or removing a sample block in this worksheet updates every plot and fit table automatically. Initial mole rows are identified by labels in each time column.",
             "Each plot contains every experiment for which its transform is defined.",
             "Measured replicate means are connected in time order using 8-point markers.",
             "Axis labels use [M] for metal concentration and 16-point type. Legends are fixed in the upper-right corner inside the axes.",
             "Untreated graphs show points without fitted lines. Only selected model windows with R² strictly above 0.99 have fitted lines.",
             "The additional first_order_early_aa_pair and aa_late_first_order_pair graphs show the two jointly optimized First-order + AA segments. Their lines are limited to the selected windows and are shown even when R² is below 0.99; the legend gives each R².",
             "Plots have no titles. Each graph spans 0–180 min (sqrt(180) for Fickian); sheets with later data also receive full-range PNGs.",
             "Vertical limits include zero and round outward to the next 0.05. Negative ln(C) values remain visible below zero.",
             "", "## Best early/late model pairs", "",
             "| Experiment | Early model | Early R² | Late model | Late R² | Weighted mean R² | Adjustment | Time relation |",
             "|---|---|---:|---|---:|---:|---|---|"]
    for pair in pairs:
        lines.append(f"| {pair['experiment']} | {pair['early_model']} | {pair['early_r2']:.5f} | "
                     f"{pair['late_model']} | {pair['late_r2']:.5f} | "
                     f"{pair['weighted_mean_r2']:.5f} | {pair['adjustment_reason'] or 'none'} | "
                     f"{pair['segment_relation']} |")
    for heading, rows, label in (
        ("Fickian early + AA/AB late", fickian_rows, "Fickian"),
        ("First-order early + AA/AB late", first_order_rows, "First-order")
    ):
        lines += ["", f"## {heading}", "",
                  f"| Experiment | Second-order model | {label} range (min) | {label} R² | Late range (min) | Late R² | Preferred | Status |",
                  "|---|---|---:|---:|---:|---:|---|---|"]
        for row in rows:
            early_range = (f"{row['early_from']:g}–{row['early_to']:g}"
                           if row["status"] == "fitted" else "unavailable")
            early_r2 = f"{row['early_r2']:.5f}" if row["status"] == "fitted" else "—"
            late_range = (f"{row['late_from']:g}–{row['late_to']:g}"
                          if row["status"] == "fitted" else "unavailable")
            late_r2 = f"{row['late_r2']:.5f}" if row["status"] == "fitted" else "—"
            lines.append(f"| {row['experiment']} | {row['late_model']} | "
                         f"{early_range} | {early_r2} | "
                         f"{late_range} | {late_r2} | "
                         f"{'yes' if row['preferred_second_order'] else 'no'} | {row['status']} |")
    for pair in pairs:
        name = pair["experiment"]
        lines += ["", f"## {name}", "",
                  f"Replicates: {next(e['replicates'].shape[1] for e in experiments if e['name']==name)}. ",
                  "| Segment | Selected model | Time (min) | R² | Apparent k | Units |",
                  "|---|---|---:|---:|---:|---|"]
        for phase in ("early", "late"):
            lines.append(f"| {phase} | {pair[f'{phase}_model']} | "
                         f"{pair[f'{phase}_from']:g}–{pair[f'{phase}_to']:g} | "
                         f"{pair[f'{phase}_r2']:.5f} | {pair[f'{phase}_k']:.6g} | "
                         f"{pair[f'{phase}_unit']} |")
        omitted_text = f"{pair['omitted_times']} min" if pair["omitted_times"] else "none"
        initial_text = (f"{pair['initial_omitted_times']} min"
                        if pair["initial_omitted_times"] else "none")
        lines += ["", f"Time relation: {pair['segment_relation']}. "
                  f"Excluded initial transient: {initial_text} ({pair['initial_omitted_points']} points). "
                  f"Omitted transition times: {omitted_text}. "
                  f"Unfitted tail after selected late window: "
                  f"{pair['late_unused_tail_times'] or 'none'} min ({pair['late_unused_tail_points']} points). "
                  f"Both selected segments R² > 0.99: {'yes' if pair['both_above_099'] else 'no'}.", "",
                  "| Model | Full R² | Early R² | Late R² |",
                  "|---|---:|---:|---:|"]
        for model in MODELS:
            if (name, model) not in fits:
                lines.append(f"| {model} | unavailable | unavailable | unavailable |")
            else:
                full, _, _, _, early, late = fits[(name, model)]
                lines.append(f"| {model} | {full['r2']:.5f} | {early['r2']:.5f} | {late['r2']:.5f} |")
    if ab_reasons:
        lines += ["", "## AB availability", ""]
        for name, reason in ab_reasons.items():
            lines.append(f"- {name}: {reason}.")
    if label_notes:
        lines += ["", "## Source metal labels", "",
                  "Numeric metal moles in the worksheet are used as recorded; verify these labels:"]
        lines.extend(f"- {note}." for note in label_notes)
    lines += ["", "## Method", "",
              "Column positions and initial-mole row numbers are discovered from worksheet labels. Missing or nonpositive initial metal moles make AB unavailable for that experiment; the other three models remain fitted.",
              "Exactly two fitted windows per model have at least six consecutive points each. The early fit can start after at most five initial transient observations; the late fit may stop before a later deflection.",
              "For a given model diagnostic or model-pair comparison, all eligible experiments are optimized together. The starting and ending observation positions of each segment differ by no more than two across samples. These are row positions within each sample, not equal clock times.",
              "There may be 0–2 omitted transition points. The remaining tail is listed, not fitted as a third segment.",
              "For every allowed early/late mechanism combination, search initial start, early endpoint, 0–2 omitted transition points, and late endpoint across all samples at once. Maximize the number of samples with both R² > 0.99; then maximize the sum of per-sample mean R² minus 0.001 per omitted initial, transition, or tail point.",
              "The early and late windows are optimized together and never share an observation. Fickian is never late; repeated mechanisms and two second-order mechanisms together are disallowed in best pairs.",
              "Restricted Fickian and first-order comparisons independently optimize their early and AA/AB late windows. Preferred marks the higher combined fit score among available second-order models.",
              "model_fits.csv contains both model-specific diagnostics and the exact selected pair rows, distinguished by fit_scope and pair_id. Graph lines show model-specific diagnostics; the pair tables record the jointly optimized windows.",
              "The two First-order + AA pair plots instead use the exact first_order_second_order pair row for AA; each line corresponds to its selected early or late window and does not extend through the transition.",
              "The point-weighted mean R² is descriptive; it is not the selection score.",
              "The AB ordinate is the replicate mean of [ln(B/A) - ln(B0/A0)]/(B0-A0), with B(t)=B0-[A0-A(t)].",
              "A high R² after a transformation does not by itself prove a physical mechanism; the Fickian sqrt-time slope is not a diffusion coefficient.", ""]
    (args.out / f"{OUTPUT_STEM}_report.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote plots, tables, and PDF report to {args.out}")
    for exp in experiments:
        print(f"{exp['name']}: {exp['replicates'].shape[1]} replicate(s); "
              f"AB {'unavailable ('+ab_reasons[exp['name']]+')' if exp['name'] in ab_reasons else 'available'}")
    for note in label_notes:
        print(f"Source label warning: {note}")


if __name__ == "__main__":
    main()
